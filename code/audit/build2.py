import json,openpyxl
from openpyxl.styles import Font,PatternFill,Alignment
rows=json.load(open('rows.json')); gi=json.load(open('gitinfo.json')); adv=json.load(open('adv_meta.json'))
GITLAB={37:('sequoia-pgp/sequoia','2024-06-26','openpgp: Fix infinite loop handling unsupported certs in raw parser.'),
 53:('sequoia-pgp/sequoia','2023-05-17','buffered-reader: Fix returning partial reads ending in errors.'),
 54:('sequoia-pgp/sequoia','2023-05-12','openpgp: Fix mapping of synthetic packets.'),
 253:('tprodanov/bam','2021-02-16','Fix unitialized memory potential errors')}
FIXADV={74:('RUSTSEC-2020-0071',['>= 0.2.23'])}
VERDICT={
 2:('WRONG','Commit is the 65-file "Replace setjmp/longjmp" refactor (2025-09), not the fix for RUSTSEC-2025-0112 (patched in 38.0.3). No usable pair.'),
 50:('WRONG','RUSTSEC-2023-0071 has NO upstream patch. SHA is the 31-file num-bigint-dig -> crypto-bigint migration (2025-02).'),
 87:('WRONG','Commit dated 2015-07, five years before the 2020 advisory; RUSTSEC-2020-0113 was never patched. Cannot be the fix.'),
 133:('WRONG','Advisory RUSTSEC-2022-0085 names fix commit 093fb5d0; this SHA is a test-only commit. Crate column also wrong (should be matrix-sdk-crypto).'),
 159:('WRONG','SHA is a 2023 performance commit ("Faster encoding for lengths in BufferBackend"). Advisory is the 2019 clone/UAF, patched in 0.6.4 / 0.7.1.'),
 213:('WRONG','Advisory RUSTSEC-2021-0051 names fix commit dd59b306 (uninitialized memory). This SHA is a 2023 merge fixing a different bug (read_length >128).'),
 216:('WRONG','Diff fixes tail_head_slice (MaybeUninit). CVE-2021-29938 / RUSTSEC-2021-0047 is the drain_filter double-drop, never patched. Verified by reading the diff.'),
 220:('WRONG','RUSTSEC-2021-0041 (big-exponent DoS) was never patched. SHA is a 2019 commit about isize subtraction edge cases.'),
 237:('WRONG','SHA is "bump v0.17.0" - Cargo.toml + Changelog only. Zero Rust delta.'),
 249:('WRONG','SHA is the release merge "Release v0.6.1" - CHANGELOG + Cargo.toml only. Real fix is the IntoIter Clone change.'),
 250:('WRONG','Advisory patched in 2.0.1 (2020). SHA is a 2025 eight-file rewrite of the UnsafeCell/queue design. Verified by reading the diff.'),
 11:('SHA NOT FOUND','Not present in ogham/rust-users. RUSTSEC-2025-0040 was never patched upstream, so no fix commit exists.'),
 175:('SHA NOT FOUND','Not present in apache/incubator-teaclave-sgx-sdk or apache/teaclave-sgx-sdk.'),
 186:('SHA NOT FOUND','Not present in hyperium/hyper. Duplicate of row 200 (same CVE), which has the correct SHA 1fb719e0. Crate column also misspelled ("hper crate").'),
 240:('SHA NOT FOUND','Repo oberien/abox no longer exists; SHA unverifiable.'),
 115:('NO RUST DELTA','Correct fix, but it only bumps Cargo.toml dependencies - no Rust pair extractable.'),
 162:('NO RUST DELTA','Correct fix, but it is JavaScript (src/theme/searcher/searcher.js) - no Rust pair.'),
 74:('REVIEW','Diff-inspected: the commit DOES remove the unsound localtime_r/tzset block, so it is the 0.2.23 fix - but it is a 12-file, 451-line "v0.3 backports" release commit, a poor extraction pair. (Advisory column corrected to RUSTSEC-2020-0071; the CVE also appears in chrono RUSTSEC-2020-0159.)'),
 207:('REVIEW','Plausible fix (decoder read path reworked) but a broad 106+/83- refactor dated before the advisory; confirm it is the soundness fix.'),
 209:('REVIEW','RUSTSEC-2021-0065 is an "unmaintained" advisory pointing at PR #32; SHA is a 2022 rename+Extend fix, never released to crates.io.'),
 212:('REVIEW','Fix is real ("Remove all unsafe code", fixes #3) but a 2025 whole-crate rewrite - a poor vulnerable/fixed pair.'),
 215:('REVIEW','Diff-inspected: content matches the advisory (insert_many size_hint hardening), but the commit predates the advisory by 536 days - confirm the parent really is the vulnerable version.'),
 168:('LABEL ERROR','SHA and advisory are correct (pyo3 reference-count bug), but Program/Crate says "futures-task crate".'),
 100:('LABEL ERROR','SHA correct. Crate column holds a repo path ("hyperium/hyper") instead of a crate name.'),
}
DIFF_READED={16,22,28,43,74,91,103,111,149,159,182,184,207,215,216,220,221,237,239,250}
DUPS={72:208,208:72,84:226,226:84,85:211,211:85,97:210,210:97,205:247,247:205,186:200,200:186}
wb=openpyxl.load_workbook('/root/.claude/uploads/a628a02a-528c-50f2-9078-58a914f12caf/1a4807f3-Gurman_Halurust_Combined_Dataset_Till_Date.xlsx')
ws=wb['Sheet1']
hdrs=['Repo (resolved)','Commit found','Commit date','Commit subject','.rs files / total','RustSec ID','Upstream patched?','How verified','VERDICT','Note']
for j,h in enumerate(hdrs,start=6):
    c=ws.cell(1,j,h); c.font=Font(name='Arial',bold=True); c.alignment=Alignment(wrap_text=True,vertical='top')
for c in ws[1][:5]: c.font=Font(name='Arial',bold=True)
FILL={'WRONG':'FFC7CE','SHA NOT FOUND':'FFC7CE','NO RUST DELTA':'FFEB9C','REVIEW':'FFEB9C','LABEL ERROR':'FFEB9C','OK':'C6EFCE'}
counts={}
for r in rows:
    rn=r['row']; g=gi.get(f"{r['repo']}@{r['sha']}") if r['repo'] else None
    aid=r['adv_ids'][0] if r['adv_ids'] else ''
    pat=adv[aid]['patched'] if aid and adv.get(aid) else None
    if rn in FIXADV: aid,pat=FIXADV[rn]
    repo=r['repo'] or (GITLAB[rn][0]+' (GitLab)' if rn in GITLAB else '')
    if g and g.get('ok'):
        found='yes'; date=g['date'][:10]; subj=g['subject'][:120]
        nrs=sum(1 for f in g['files'] if f['f'].endswith('.rs')); files=f"{nrs}/{g['nfiles']}"; nf=g['nfiles']
    elif rn in GITLAB:
        found='yes'; date=GITLAB[rn][1]; subj=GITLAB[rn][2]; files='n/a'; nf=0
    else:
        found='NO'; date=''; subj=''; files=''; nf=0
    v,note=VERDICT.get(rn,('OK',''))
    if v=='OK' and pat==[]:
        v,note='REVIEW','Advisory has no upstream patch (patched = []); the "safe" side is a later voluntary cleanup, not a released fix.'
    if v=='OK' and not aid:
        note='No RustSec advisory for this CVE (application-level project) - verified only that the SHA exists and its subject describes the fix.'
    if nf>20 and v=='OK':
        note=(note+' ' if note else '')+f'Large diff ({nf} files) - extraction needs care.'
    if rn in DUPS: note=(note+' ' if note else '')+f'Duplicate CVE row (see row {DUPS[rn]}).'
    how='diff inspected' if rn in DIFF_READED else ('advisory matched' if aid else 'subject only')
    if found=='NO': how='could not verify'
    ws.cell(rn,6,repo); ws.cell(rn,7,found); ws.cell(rn,8,date); ws.cell(rn,9,subj)
    ws.cell(rn,10,files); ws.cell(rn,11,aid); ws.cell(rn,12,'yes' if pat else ('NO' if pat==[] else '?')); ws.cell(rn,13,how)
    c=ws.cell(rn,14,v); c.fill=PatternFill('solid',fgColor=FILL.get(v,'FFFFFF')); c.font=Font(name='Arial',bold=True)
    ws.cell(rn,15,note)
    counts[v]=counts.get(v,0)+1
    for j in range(1,16):
        cc=ws.cell(rn,j)
        if cc.font.name!='Arial': cc.font=Font(name='Arial')
        cc.alignment=Alignment(wrap_text=True,vertical='top')
for col,w in zip(['A','B','C','D','E','F','G','H','I','J','K','L','M','N','O'],[16,11,20,18,42,26,9,11,46,10,18,10,16,15,72]):
    ws.column_dimensions[col].width=w
ws.freeze_panes='A2'; ws.auto_filter.ref=f'A1:O{ws.max_row}'
wb.save('/tmp/out/Halurust_SHA_audit.xlsx')
print(counts)
import collections
print(collections.Counter(ws.cell(r,13).value for r in range(2,258) if ws.cell(r,1).value))
