import openpyxl, os, re, json, glob
XL='/root/.claude/uploads/a628a02a-528c-50f2-9078-58a914f12caf/1a4807f3-Gurman_Halurust_Combined_Dataset_Till_Date.xlsx'
wb=openpyxl.load_workbook(XL,data_only=True); ws=wb['Sheet1']
rows=[]
for r in range(2,ws.max_row+1):
    v=[ws.cell(r,c).value for c in range(1,6)]
    if all(x is None for x in v): continue
    rows.append(dict(row=r,cve=(v[0] or '').strip(),cwe=(v[1] or '').strip(),crate=(v[2] or '').strip(),
                     version=str(v[3] or '').strip(),sha=str(v[4] or '').strip()))
print('data rows:',len(rows))
bad=[x for x in rows if not re.fullmatch(r'[0-9a-f]{40}',x['sha'] or '')]
print('non-40hex sha rows:',len(bad))
for b in bad[:20]: print('  ',b['row'],b['cve'],b['crate'],repr(b['sha']))
from collections import Counter
c=Counter(x['cve'] for x in rows)
dups={k:v for k,v in c.items() if v>1}
print('duplicate CVE ids:',len(dups),dups)
cs=Counter(x['sha'] for x in rows)
dsha={k:v for k,v in cs.items() if v>1}
print('duplicate SHAs:',len(dsha))
for k,v in list(dsha.items())[:20]:
    print('  ',k,v,[ (x['row'],x['cve'],x['crate']) for x in rows if x['sha']==k])
json.dump(rows,open('/tmp/hr/rows.json','w'))
